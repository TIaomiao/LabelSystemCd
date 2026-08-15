#!/usr/bin/env python3
"""Build the physician-facing CMR Workstation manual from its Markdown source."""

from __future__ import annotations

import re
from pathlib import Path
from typing import Iterable

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


WORKSPACE = Path(__file__).resolve().parents[1]
SOURCE = WORKSPACE / "CMR工作站医生操作说明书.md"
OUTPUT = WORKSPACE / "CMR工作站医生操作说明书_医生版.docx"
ASSET_DIR = WORKSPACE / "manual_assets"
LAYOUT_IMAGE = ASSET_DIR / "工作台布局示意.png"
FLOW_IMAGE = ASSET_DIR / "最小上手流程.png"

FONT_REGULAR = Path("/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc")
FONT_BOLD = Path("/usr/share/fonts/opentype/noto/NotoSansCJK-Bold.ttc")
FONT_NAME = "Noto Sans CJK SC"

COLORS = {
    "ink": "202629",
    "panel": "2A3236",
    "text": "202629",
    "muted": "68757A",
    "line": "D7DDDF",
    "soft": "F3F5F5",
    "green": "A8CB4E",
    "green_dark": "6F8736",
    "green_soft": "EEF5DA",
    "cyan": "22A2C3",
    "cyan_soft": "E7F5F8",
    "orange": "D28A35",
    "orange_soft": "FBF0E2",
    "white": "FFFFFF",
}


def font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    return ImageFont.truetype(str(FONT_BOLD if bold else FONT_REGULAR), size=size)


def rounded_box(draw: ImageDraw.ImageDraw, xy, fill, outline=None, radius=12, width=2):
    draw.rounded_rectangle(xy, radius=radius, fill=fill, outline=outline, width=width)


def center_text(draw: ImageDraw.ImageDraw, box, text: str, text_font, fill: str, spacing: int = 5):
    left, top, right, bottom = box
    wrapped = text.split("\n")
    heights = []
    for line in wrapped:
        bounds = draw.textbbox((0, 0), line, font=text_font)
        heights.append(bounds[3] - bounds[1])
    total = sum(heights) + spacing * max(0, len(wrapped) - 1)
    y = top + (bottom - top - total) / 2
    for line, height in zip(wrapped, heights):
        bounds = draw.textbbox((0, 0), line, font=text_font)
        width = bounds[2] - bounds[0]
        draw.text((left + (right - left - width) / 2, y), line, font=text_font, fill=fill)
        y += height + spacing


def build_layout_image(path: Path) -> None:
    image = Image.new("RGB", (1800, 760), "#FFFFFF")
    draw = ImageDraw.Draw(image)
    draw.rectangle((0, 0, 1800, 10), fill="#A8CB4E")
    draw.text((62, 38), "CMR Workstation", font=font(42, True), fill="#202629")
    draw.text((425, 50), "功能模块", font=font(25, True), fill="#68757A")
    draw.text((1120, 50), "当前病例", font=font(25, True), fill="#68757A")
    rounded_box(draw, (1475, 27, 1716, 91), fill="#A8CB4E", radius=8)
    center_text(draw, (1475, 27, 1716, 91), "AI 专家", font(28, True), "#172025")
    draw.line((42, 118, 1758, 118), fill="#D7DDDF", width=2)

    regions = [
        ((42, 154, 430, 708), "1  病例库", "选择数据来源\n搜索登记号/病例号\n导入或打开病例", "#F3F5F5", "#A8CB4E"),
        ((456, 154, 1342, 708), "2  主工作区", "影像浏览与勾画\n评估、报告评分、质控\nSlice / Phase Matrix", "#FFFFFF", "#22A2C3"),
        ((1368, 154, 1758, 708), "3  工作站专家", "操作指导\n截图反馈\n问题单沉淀", "#F3F5F5", "#D28A35"),
    ]
    for box, title, body, fill_color, accent in regions:
        rounded_box(draw, box, fill=fill_color, outline="#D7DDDF", radius=10, width=2)
        draw.rectangle((box[0], box[1], box[0] + 10, box[3]), fill=accent)
        draw.text((box[0] + 38, box[1] + 42), title, font=font(31, True), fill="#202629")
        draw.multiline_text((box[0] + 38, box[1] + 132), body, font=font(25), fill="#68757A", spacing=18)
    image.save(path, quality=95)


def build_flow_image(path: Path) -> None:
    image = Image.new("RGB", (1800, 430), "#FFFFFF")
    draw = ImageDraw.Draw(image)
    draw.text((40, 24), "一次完整操作的最短路径", font=font(34, True), fill="#202629")
    steps = [
        ("01", "选择病例库", "搜索并核对病例"),
        ("02", "导入 / 打开", "等待影像加载"),
        ("03", "核对序列", "确认 slice / phase"),
        ("04", "完成勾画", "选 ROI 后人工复核"),
        ("05", "保存结果", "等待成功提示"),
        ("06", "确认与反馈", "看矩阵或询问 AI"),
    ]
    left = 42
    box_w = 260
    gap = 32
    top, bottom = 118, 366
    for index, (number, title, body) in enumerate(steps):
        x1 = left + index * (box_w + gap)
        x2 = x1 + box_w
        rounded_box(draw, (x1, top, x2, bottom), fill="#F3F5F5", outline="#D7DDDF", radius=10, width=2)
        draw.rectangle((x1, top, x2, top + 12), fill="#A8CB4E" if index < 5 else "#22A2C3")
        draw.text((x1 + 22, top + 32), number, font=font(25, True), fill="#6F8736")
        draw.text((x1 + 22, top + 85), title, font=font(27, True), fill="#202629")
        draw.multiline_text((x1 + 22, top + 137), body, font=font(21), fill="#68757A", spacing=8)
        if index < len(steps) - 1:
            arrow_y = (top + bottom) // 2
            arrow_x1 = x2 + 7
            arrow_x2 = x2 + gap - 7
            draw.line((arrow_x1, arrow_y, arrow_x2, arrow_y), fill="#8D9A9F", width=5)
            draw.polygon([(arrow_x2, arrow_y), (arrow_x2 - 13, arrow_y - 9), (arrow_x2 - 13, arrow_y + 9)], fill="#8D9A9F")
    image.save(path, quality=95)


def set_cell_shading(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def set_cell_border(cell, **edges) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge_name, values in edges.items():
        edge = borders.find(qn(f"w:{edge_name}"))
        if edge is None:
            edge = OxmlElement(f"w:{edge_name}")
            borders.append(edge)
        for key, value in values.items():
            edge.set(qn(f"w:{key}"), str(value))


def set_cell_margins(cell, top=80, start=100, bottom=80, end=100) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        element = tc_mar.find(qn(f"w:{margin}"))
        if element is None:
            element = OxmlElement(f"w:{margin}")
            tc_mar.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def set_run_font(run, size: float | None = None, bold: bool | None = None, color: str | None = None) -> None:
    run.font.name = FONT_NAME
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_inline_markdown(paragraph, text: str, size: float | None = None, color: str | None = None) -> None:
    pattern = re.compile(r"(\*\*.+?\*\*|`.+?`)")
    cursor = 0
    for match in pattern.finditer(text):
        if match.start() > cursor:
            run = paragraph.add_run(text[cursor:match.start()])
            set_run_font(run, size=size, color=color)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, size=size, bold=True, color=color)
        else:
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, size=(size or 10.5) - 0.5, color=COLORS["green_dark"])
            run.font.name = "Consolas"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)
        cursor = match.end()
    if cursor < len(text):
        run = paragraph.add_run(text[cursor:])
        set_run_font(run, size=size, color=color)


def add_bottom_border(paragraph, color: str, size: int = 10) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "5")
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    set_run_font(run, size=8.5, color=COLORS["muted"])
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])
    run2 = paragraph.add_run(" 页")
    set_run_font(run2, size=8.5, color=COLORS["muted"])


def configure_styles(document: Document) -> None:
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = FONT_NAME
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(COLORS["text"])
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.28

    for name, size, color, before, after in (
        ("Title", 29, COLORS["ink"], 0, 0),
        ("Heading 1", 18, COLORS["ink"], 18, 8),
        ("Heading 2", 14, COLORS["green_dark"], 14, 6),
        ("Heading 3", 11.5, COLORS["ink"], 10, 4),
    ):
        style = styles[name]
        style.font.name = FONT_NAME
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def configure_document_settings(document: Document) -> None:
    settings = document.settings._element
    compat = settings.find(qn("w:compat"))
    if compat is None:
        compat = OxmlElement("w:compat")
        settings.append(compat)
    compatibility_mode = next(
        (
            item
            for item in compat.findall(qn("w:compatSetting"))
            if item.get(qn("w:name")) == "compatibilityMode"
        ),
        None,
    )
    if compatibility_mode is None:
        compatibility_mode = OxmlElement("w:compatSetting")
        compatibility_mode.set(qn("w:name"), "compatibilityMode")
        compatibility_mode.set(qn("w:uri"), "http://schemas.microsoft.com/office/word")
        compat.append(compatibility_mode)
    compatibility_mode.set(qn("w:val"), "15")


def configure_sections(document: Document) -> None:
    for index, section in enumerate(document.sections):
        section.page_width = Cm(21)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(1.55)
        section.bottom_margin = Cm(1.55)
        section.left_margin = Cm(1.75)
        section.right_margin = Cm(1.75)
        section.header_distance = Cm(0.7)
        section.footer_distance = Cm(0.7)

        header = section.header
        footer = section.footer
        if index == 0:
            header.is_linked_to_previous = False
            footer.is_linked_to_previous = False
            header.paragraphs[0].clear()
            footer.paragraphs[0].clear()
            continue

        header.is_linked_to_previous = False
        header_p = header.paragraphs[0]
        header_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        header_p.clear()
        run = header_p.add_run("CMR WORKSTATION  /  医生操作说明书")
        set_run_font(run, size=8.5, bold=True, color=COLORS["muted"])
        add_bottom_border(header_p, COLORS["line"], 5)

        footer.is_linked_to_previous = False
        footer_p = footer.paragraphs[0]
        footer_p.clear()
        add_page_number(footer_p)
        if index == 1:
            pg_num_type = OxmlElement("w:pgNumType")
            pg_num_type.set(qn("w:start"), "1")
            section._sectPr.append(pg_num_type)


def add_cover(document: Document) -> None:
    p = document.add_paragraph()
    p.paragraph_format.space_before = Pt(36)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run("CMR / CLINICAL WORKSTATION")
    set_run_font(r, size=11, bold=True, color=COLORS["green_dark"])
    p.paragraph_format.space_after = Pt(20)

    p = document.add_paragraph()
    r = p.add_run("医生操作说明书")
    set_run_font(r, size=32, bold=True, color=COLORS["ink"])
    p.paragraph_format.space_after = Pt(12)

    p = document.add_paragraph()
    r = p.add_run("从 3 分钟上手，到影像勾画、AI 辅助、结果复核与问题反馈")
    set_run_font(r, size=12.5, color=COLORS["muted"])
    p.paragraph_format.space_after = Pt(22)
    add_bottom_border(p, COLORS["green"], 16)

    p = document.add_paragraph()
    p.paragraph_format.space_before = Pt(20)
    r = p.add_run("适用版本")
    set_run_font(r, size=9.5, bold=True, color=COLORS["muted"])
    r = p.add_run("   v2026.08.15-2")
    set_run_font(r, size=11, bold=True, color=COLORS["ink"])

    p = document.add_paragraph()
    r = p.add_run("更新日期   2026-08-15")
    set_run_font(r, size=10, color=COLORS["text"])

    p = document.add_paragraph()
    r = p.add_run("适用对象   标注、复核、报告评分与科研质控医生及实习生")
    set_run_font(r, size=10, color=COLORS["text"])

    contact = document.add_paragraph()
    contact.paragraph_format.space_before = Pt(150)
    r = contact.add_run("作者 / 工作站维护联系人")
    set_run_font(r, size=9.5, bold=True, color=COLORS["green_dark"])

    contact = document.add_paragraph()
    r = contact.add_run("Zian")
    set_run_font(r, size=15, bold=True, color=COLORS["ink"])
    contact.paragraph_format.space_after = Pt(2)

    contact = document.add_paragraph()
    r = contact.add_run("1980651739@qq.com")
    set_run_font(r, size=10.5, color=COLORS["muted"])

    note = document.add_paragraph()
    note.paragraph_format.space_before = Pt(38)
    add_inline_markdown(note, "AI 分割、传播、测量和问答结果均需医生结合原始影像复核。", size=9.5, color=COLORS["muted"])
    document.add_section(WD_SECTION.NEW_PAGE)


def add_quick_directory(document: Document) -> None:
    title = document.add_paragraph(style="Heading 1")
    title.add_run("快速目录")
    add_bottom_border(title, COLORS["green"], 14)

    intro = document.add_paragraph()
    add_inline_markdown(intro, "第一次使用先读左列；遇到具体任务时再查右列。", color=COLORS["muted"])

    table = document.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    left, right = table.rows[0].cells
    for cell in (left, right):
        set_cell_margins(cell, top=200, start=220, bottom=200, end=220)
        set_cell_border(cell, top={"val": "single", "sz": 8, "color": COLORS["line"]}, bottom={"val": "single", "sz": 8, "color": COLORS["line"]}, start={"val": "single", "sz": 8, "color": COLORS["line"]}, end={"val": "single", "sz": 8, "color": COLORS["line"]})
    set_cell_shading(left, COLORS["green_soft"])
    set_cell_shading(right, COLORS["cyan_soft"])

    p = left.paragraphs[0]
    add_inline_markdown(p, "第一次使用", size=13, color=COLORS["green_dark"])
    p.runs[0].bold = True
    for item in (
        "3 分钟最小上手指南",
        "工作台界面速查",
        "病例选择、导入与切换",
        "离开病例前的 20 秒检查",
    ):
        p = left.add_paragraph(style="List Bullet")
        add_inline_markdown(p, item, size=10)

    p = right.paragraphs[0]
    add_inline_markdown(p, "按需查阅", size=13, color=COLORS["cyan"])
    p.runs[0].bold = True
    for item in (
        "CMR 影像浏览与勾画",
        "AI 分割与传播补全",
        "2D 应变与追踪验证",
        "AI 专家与截图反馈",
        "各评估、报告、质控模块",
        "常见问题与术语速查",
    ):
        p = right.add_paragraph(style="List Bullet")
        add_inline_markdown(p, item, size=10)

    document.add_paragraph()
    callout = document.add_table(rows=1, cols=1)
    callout.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = callout.cell(0, 0)
    set_cell_shading(cell, COLORS["orange_soft"])
    set_cell_margins(cell, top=150, start=180, bottom=150, end=180)
    set_cell_border(cell, start={"val": "single", "sz": 22, "color": COLORS["orange"]}, top={"val": "nil"}, bottom={"val": "nil"}, end={"val": "nil"})
    p = cell.paragraphs[0]
    run = p.add_run("高影响操作提醒  ")
    set_run_font(run, size=10.5, bold=True, color=COLORS["orange"])
    add_inline_markdown(p, "传播补全、清空当前帧全部轮廓和批量 AI 会改变标注结果；执行前先核对病例、序列、slice 和 phase。", size=10.5)

    picture = document.add_picture(str(LAYOUT_IMAGE), width=Cm(17.2))
    picture._inline.graphic.graphicData.pic.spPr.xfrm.ext.cx = picture._inline.extent.cx
    picture._inline.graphic.graphicData.pic.spPr.xfrm.ext.cy = picture._inline.extent.cy
    document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption = document.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_inline_markdown(caption, "工作台布局：病例库、主工作区与 AI 专家。", size=8.8, color=COLORS["muted"])


def add_callout(document: Document, text: str, kind: str = "tip") -> None:
    palette = {
        "tip": (COLORS["green_soft"], COLORS["green_dark"], "提示"),
        "warning": (COLORS["orange_soft"], COLORS["orange"], "注意"),
        "info": (COLORS["cyan_soft"], COLORS["cyan"], "说明"),
    }
    fill, accent, label = palette[kind]
    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_margins(cell, top=120, start=160, bottom=120, end=160)
    set_cell_border(cell, start={"val": "single", "sz": 18, "color": accent}, top={"val": "nil"}, bottom={"val": "nil"}, end={"val": "nil"})
    p = cell.paragraphs[0]
    run = p.add_run(f"{label}  ")
    set_run_font(run, size=10, bold=True, color=accent)
    add_inline_markdown(p, text, size=10)
    document.add_paragraph().paragraph_format.space_after = Pt(1)


def add_markdown_table(document: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    column_count = max(len(row) for row in rows)
    table = document.add_table(rows=len(rows), cols=column_count)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    table.style = "Table Grid"
    for row_index, values in enumerate(rows):
        for col_index in range(column_count):
            cell = table.cell(row_index, col_index)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)
            value = values[col_index] if col_index < len(values) else ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            add_inline_markdown(p, value, size=8.7 if column_count >= 4 else 9.2)
            if row_index == 0:
                set_cell_shading(cell, COLORS["panel"])
                for run in p.runs:
                    run.bold = True
                    run.font.color.rgb = RGBColor.from_string(COLORS["white"])
            elif row_index % 2 == 0:
                set_cell_shading(cell, COLORS["soft"])
            if row_index > 0 and col_index == 0:
                for run in p.runs:
                    run.bold = True
                    run.font.color.rgb = RGBColor.from_string(COLORS["green_dark"])
    document.add_paragraph().paragraph_format.space_after = Pt(1)


def parse_table(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    raw = []
    index = start
    while index < len(lines) and lines[index].strip().startswith("|"):
        raw.append(lines[index].strip())
        index += 1
    rows = []
    for line in raw:
        values = [value.strip() for value in line.strip("|").split("|")]
        if all(re.fullmatch(r":?-{3,}:?", value.replace(" ", "")) for value in values):
            continue
        rows.append(values)
    return rows, index


def parse_markdown(document: Document, markdown: str) -> None:
    lines = markdown.splitlines()
    start = next(index for index, line in enumerate(lines) if line.startswith("## 先看这里"))
    index = start
    paragraph_buffer: list[str] = []
    def flush_paragraph() -> None:
        if not paragraph_buffer:
            return
        text = "".join(part.strip() for part in paragraph_buffer).strip()
        paragraph_buffer.clear()
        if not text:
            return
        p = document.add_paragraph()
        add_inline_markdown(p, text)

    while index < len(lines):
        raw = lines[index]
        line = raw.strip()

        if not line or line == "---":
            flush_paragraph()
            index += 1
            continue

        if line.startswith("|"):
            flush_paragraph()
            rows, index = parse_table(lines, index)
            add_markdown_table(document, rows)
            continue

        heading = re.match(r"^(#{2,4})\s+(.+)$", line)
        if heading:
            flush_paragraph()
            level = len(heading.group(1))
            title = heading.group(2).strip()
            p = document.add_paragraph(style=f"Heading {level - 1}")
            add_inline_markdown(p, title)
            if level == 2:
                add_bottom_border(p, COLORS["green"], 10)
            if title.startswith("先看这里"):
                document.add_picture(str(FLOW_IMAGE), width=Cm(17.2))
                document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            if title == "5. AI 分割与传播补全":
                add_callout(document, "AI 和传播用于生成待复核结果，不代表医生已经确认。任务运行中不要重复点击。", "warning")
            if title.startswith("7. AI 专家"):
                add_callout(document, "AI 专家会提供操作指导并沉淀问题单，但不会直接修改病例、标注、权限或服务器。", "info")
            index += 1
            continue

        if line.startswith(">"):
            flush_paragraph()
            quote_lines: list[str] = []
            while index < len(lines) and lines[index].strip().startswith(">"):
                quote_lines.append(lines[index].strip().lstrip("> ").strip())
                index += 1
            quote = "\n".join(part for part in quote_lines if part)
            kind = "warning" if any(token in quote for token in ("不要", "必须", "未发现 DICOM")) else "tip"
            add_callout(document, quote, kind)
            continue

        numbered = re.match(r"^(\d+)\.\s+(.+)$", line)
        if numbered:
            flush_paragraph()
            p = document.add_paragraph(style="List Number")
            add_inline_markdown(p, numbered.group(2))
            index += 1
            continue

        checkbox = re.match(r"^-\s+\[\s*\]\s+(.+)$", line)
        if checkbox:
            flush_paragraph()
            p = document.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.45)
            p.paragraph_format.first_line_indent = Cm(-0.45)
            run = p.add_run("□  ")
            set_run_font(run, size=11, bold=True, color=COLORS["green_dark"])
            add_inline_markdown(p, checkbox.group(1))
            index += 1
            continue

        bullet = re.match(r"^-\s+(.+)$", line)
        if bullet:
            flush_paragraph()
            p = document.add_paragraph(style="List Bullet")
            add_inline_markdown(p, bullet.group(1))
            index += 1
            continue

        paragraph_buffer.append(raw)
        index += 1

    flush_paragraph()


def add_document_properties(document: Document) -> None:
    props = document.core_properties
    props.title = "CMR Workstation 医生操作说明书"
    props.subject = "CMR 影像浏览、勾画、AI 辅助、结果复核与问题反馈"
    props.author = "Zian"
    props.keywords = "CMR, 医生操作说明书, 4CH, SAX, LGE, 心肌应变, AI 专家"
    props.comments = "基于 v2026.08.15-2 生成；工作站维护联系人：Zian"


def validate_document(document: Document) -> None:
    paragraph_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    required = (
        "3 分钟最小上手指南",
        "AI 分割与传播补全",
        "2D 应变与追踪验证",
        "AI 专家：操作说明书 Plus 与问题反馈入口",
        "常见问题与处理",
        "离开病例前的 20 秒检查",
        "1980651739@qq.com",
    )
    missing = [item for item in required if item not in paragraph_text]
    if missing:
        raise RuntimeError(f"Generated document is missing required sections: {missing}")
    if len(document.tables) < 8:
        raise RuntimeError("Generated document contains too few structured tables")
    if len(document.inline_shapes) < 2:
        raise RuntimeError("Generated document is missing visual guides")


def main() -> None:
    if not SOURCE.is_file():
        raise FileNotFoundError(SOURCE)
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    build_layout_image(LAYOUT_IMAGE)
    build_flow_image(FLOW_IMAGE)

    document = Document()
    configure_document_settings(document)
    configure_styles(document)
    add_document_properties(document)
    add_cover(document)
    add_quick_directory(document)
    parse_markdown(document, SOURCE.read_text(encoding="utf-8"))
    configure_sections(document)
    validate_document(document)
    document.save(OUTPUT)
    print(f"created: {OUTPUT}")
    print(f"paragraphs: {len(document.paragraphs)}")
    print(f"tables: {len(document.tables)}")
    print(f"visuals: {len(document.inline_shapes)}")


if __name__ == "__main__":
    main()
